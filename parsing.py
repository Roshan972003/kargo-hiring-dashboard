from io import BytesIO


class ParseError(Exception):
    pass


def extract_text(file_bytes, filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext == "docx":
        return _extract_docx(file_bytes)
    if ext == "txt":
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="ignore")

    raise ParseError(
        f"Unsupported file type '.{ext}'. Upload a PDF, DOCX, or TXT resume."
    )


def _extract_pdf(file_bytes):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(file_bytes))
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(parts)


def _extract_docx(file_bytes):
    from docx import Document

    doc = Document(BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)
